# -*- coding: utf-8 -*-
"""
강의안 레이아웃 카탈로그 (파라미터화).
각 함수는 Deck 인스턴스를 받아 슬라이드 1장을 추가한다.
색을 지정하지 않으면 코랄/틸/앰버를 자동 순환한다.

사용:
    from builder import Deck
    import layouts as L
    d = Deck(brand="9WAY · 미래인재연구소")
    L.cover(d, title="...", subtitle="...")
    ...
    d.save("out.pptx")
"""
import tokens as T
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from helpers import (px, pt, rgb, rect, oval, text, bullets, line_div, vline,
                     diag, spike_mark, add_table, picture, _apply_run)

LX = 80
CW = T.SLIDE_W_PX - LX * 2
SW = T.SLIDE_W_PX
ACCENTS = ["primary", "accent_teal", "accent_amber"]


def _accent(i, color=None):
    return color if color else ACCENTS[i % len(ACCENTS)]


def _header(slide, eyebrow, title, title_token="display_md",
            title_weight="SemiBold", y=58):
    if eyebrow:
        text(slide, LX, y, CW, 22, eyebrow.upper(), type_token="caption_upper",
             color="muted")
    text(slide, LX, y + 26, CW, 64, title, type_token=title_token,
         color="ink", weight=title_weight)


def _chip(slide, x, y, label, active=False, size=13):
    w = 30 + len(label) * (size * 0.98)
    rect(slide, x, y, w, 32, fill="primary" if active else "surface_card",
         radius=999)
    text(slide, x, y, w, 32, label, type_token="caption",
         color="on_primary" if active else "ink", align="center",
         valign="middle", size_px=size)
    return w


# ===========================================================================
# 표지 / 섹션 / 메시지
# ===========================================================================
def cover(deck, title, subtitle="", eyebrow="", brand="", title2=None,
          title2_color="primary", meta=""):
    """표지. title2를 주면 2번째 줄을 강조색으로."""
    s = deck.add("canvas")
    if brand:
        text(s, LX, 56, 400, 30, brand, type_token="title_lg", color="ink",
             weight="ExtraBold", valign="middle")
    if eyebrow:
        text(s, SW - LX - 420, 60, 420, 20, eyebrow.upper(),
             type_token="caption_upper", color="muted", align="right")
    rect(s, LX, 256, 64, 6, fill="primary", radius=3)
    if title2:
        text(s, LX, 278, 1040, 90, title, type_token="display_xl",
             color="ink", weight="Light")
        text(s, LX, 372, 1040, 90, title2, type_token="display_xl",
             color=title2_color, weight="ExtraBold")
    else:
        text(s, LX, 300, 1040, 150, title, type_token="display_xl",
             color="ink", weight="Light", line_spacing=1.08)
    if subtitle:
        text(s, LX, 502, 900, 40, subtitle, type_token="display_sm",
             color="muted", weight="Light")
    if meta:
        text(s, LX, 662, CW, 20, meta, type_token="body_sm", color="muted_soft")
    return s


def section_divider(deck, number, title, subtitle="", dark=True):
    """섹션 전환. 큰 번호 + 제목."""
    s = deck.add("surface_dark" if dark else "surface_cream_strong")
    tc = "on_dark" if dark else "ink"
    sc = "on_dark_soft" if dark else "muted"
    if number:
        text(s, LX, 150, 600, 300, str(number), type_token="display_xl",
             color="primary", weight="Black", size_px=240)
        tx, ty = 560, 300
    else:
        tx, ty = LX, 320
    rect(s, tx, ty - 22 if number else 288, 64, 6, fill="primary", radius=3)
    text(s, tx, ty + 8 if number else 332, 680, 70, title,
         type_token="display_lg", color=tc, weight="Light")
    if subtitle:
        text(s, tx, (ty + 88) if number else 432, 680, 40, subtitle,
             type_token="display_sm", color=sc, weight="Light")
    deck.footer(s, dark=dark)
    return s


def statement(deck, lines, bg="surface_cream_strong"):
    """큰 중앙 메시지 디바이더."""
    s = deck.add(bg)
    dark = bg == "surface_dark"
    for i in range(5):
        oval(s, SW / 2 - 44 + i * 22, 196, 10, 10, fill="primary")
    text(s, LX, 270, CW, 200, lines, type_token="display_lg",
         color="on_dark" if dark else "ink", align="center",
         line_spacing=1.25, valign="middle")
    line_div(s, SW / 2 - 120, 520, 240, color="primary", weight=2)
    deck.footer(s, dark=dark)
    return s


def quote(deck, quote_text, author="", bg="surface_cream_strong"):
    """인용/명언."""
    s = deck.add(bg)
    text(s, LX, 150, 200, 120, "“", type_token="display_xl",
         color="primary", weight="Black", size_px=150)
    text(s, LX, 290, 1040, 180, quote_text, type_token="display_lg",
         color="ink", weight="Light", line_spacing=1.3)
    if author:
        text(s, LX, 500, 700, 30, author, type_token="title_md", color="muted")
    deck.footer(s)
    return s


# ===========================================================================
# 학습 / 개념
# ===========================================================================
def objectives(deck, title, goals, eyebrow="Learning Objectives"):
    """체크리스트형 학습 목표. goals: [str]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    y = 200
    for g in goals:
        rect(s, LX, y, 36, 36, fill="primary", radius=8)
        text(s, LX, y, 36, 36, "✓", type_token="title_md",
             color="on_primary", align="center", valign="middle")
        text(s, LX + 56, y, 980, 36, g, type_token="title_md", color="ink",
             valign="middle", weight="Regular")
        y += 70
    text(s, 900, 200, 300, 320, str(len(goals)), type_token="display_xl",
         color="surface_card", weight="Black", size_px=300, align="right")
    deck.footer(s)
    return s


def concept(deck, term, definition, formula="", example="",
            eyebrow="Key Concept"):
    """핵심 개념: 큰 용어 + 정의 + (공식) + (예시 다크 콜아웃)."""
    s = deck.add("canvas")
    text(s, LX, 90, 200, 24, eyebrow.upper(), type_token="caption_upper",
         color="primary")
    text(s, LX, 130, 700, 120, term, type_token="display_xl", color="ink",
         weight="Black", size_px=104)
    text(s, LX, 270, 760, 110, definition, type_token="title_lg", color="body",
         weight="Regular", line_spacing=1.5)
    if formula:
        rect(s, LX, 400, 760, 90, fill="surface_card", radius=12)
        text(s, LX, 400, 760, 90, formula, type_token="display_sm",
             color="ink", align="center", valign="middle", weight="SemiBold")
    if example:
        rect(s, 880, 130, CW - (880 - LX), 360, fill="surface_dark", radius=12)
        text(s, 912, 162, 280, 24, "예시", type_token="caption_upper",
             color="accent_amber")
        text(s, 912, 196, 280, 280, example, type_token="body_md",
             color="on_dark_soft", line_spacing=1.6)
    deck.footer(s)
    return s


def question(deck, question_text, sub=""):
    """풀코랄 빅 퀘스천."""
    s = deck.add("primary")
    text(s, LX, 110, 200, 100, "?", type_token="display_xl",
         color="on_primary", weight="Black", size_px=140)
    text(s, LX, 280, 1040, 220, question_text, type_token="display_xl",
         color="on_primary", weight="Light", line_spacing=1.15, size_px=64)
    if sub:
        text(s, LX, 630, 800, 30, sub, type_token="title_md",
             color="on_primary", weight="Regular")
    return s


def takeaways(deck, title, items, eyebrow="Key Takeaways"):
    """핵심 N가지. items: [(title, desc)] (최대 4)."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = min(len(items), 4)
    cw = (CW - (n - 1) * 24) / n
    for i in range(n):
        t, desc = items[i]
        color = _accent(i)
        x = LX + i * (cw + 24)
        text(s, x, 210, cw, 90, f"{i+1:02d}", type_token="display_xl",
             color=color, weight="ExtraBold", size_px=60)
        line_div(s, x, 300, 48, color=color, weight=3)
        text(s, x, 324, cw, 40, t, type_token="title_lg", color="ink")
        text(s, x, 372, cw, 150, desc, type_token="body_md", color="muted",
             line_spacing=1.55)
    deck.footer(s)
    return s


# ===========================================================================
# 카드 / 통계
# ===========================================================================
def three_up(deck, title, cards, eyebrow=""):
    """3-up 피처 카드. cards: [(title, body[, color])]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    cw = (CW - 48) / 3
    for i, c in enumerate(cards[:3]):
        t, body = c[0], c[1]
        color = _accent(i, c[2] if len(c) > 2 else None)
        x = LX + i * (cw + 24)
        rect(s, x, 200, cw, 300, fill="surface_card", radius=12)
        oval(s, x + 28, 228, 48, 48, fill=color)
        text(s, x + 28, 294, cw - 56, 56, t, type_token="title_lg",
             color="ink", line_spacing=1.2)
        text(s, x + 28, 360, cw - 56, 130, body, type_token="body_md",
             color="body", line_spacing=1.55)
    deck.footer(s)
    return s


def stat_cards(deck, title, stats, eyebrow=""):
    """통계 카드. stats: [(number, label, desc[, color])]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = min(len(stats), 4)
    cw = (CW - (n - 1) * 24) / n
    for i in range(n):
        st = stats[i]
        num, label, desc = st[0], st[1], st[2]
        color = _accent(i, st[3] if len(st) > 3 else None)
        x = LX + i * (cw + 24)
        rect(s, x, 210, cw, 320, fill="surface_card", radius=12)
        text(s, x + 28, 238, cw - 56, 70, num, type_token="display_lg",
             color=color, size_px=52, weight="ExtraBold")
        text(s, x + 28, 320, cw - 56, 30, label, type_token="title_md",
             color="ink")
        line_div(s, x + 28, 356, 40, color=color, weight=3)
        text(s, x + 28, 372, cw - 56, 140, desc, type_token="body_md",
             color="muted", line_spacing=1.5)
    deck.footer(s)
    return s


def big_stat(deck, number, pre="", post="", note=""):
    """단일 빅 통계."""
    s = deck.add("canvas")
    if pre:
        text(s, LX, 160, CW, 40, pre, type_token="display_sm", color="muted",
             align="center", weight="Light")
    text(s, LX, 220, CW, 200, number, type_token="display_xl", color="primary",
         align="center", weight="Black", size_px=190)
    if post:
        text(s, LX, 450, CW, 40, post, type_token="display_sm", color="ink",
             align="center")
    if note:
        text(s, LX, 510, CW, 24, note, type_token="body_sm",
             color="muted_soft", align="center")
    deck.footer(s)
    return s


# ===========================================================================
# 표 / 프로세스 / 타임라인
# ===========================================================================
def comparison_table(deck, title, data, highlight_col=None, eyebrow="",
                     header_color="dark", label_col=True, intro=""):
    """비교 표(네이티브). data: 2D 리스트(첫 행=헤더).
    header_color: 'dark' | 'coral'. highlight_col: 강조 열 인덱스."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    y = 200
    if intro:
        text(s, LX, 144, CW, 36, intro, type_token="body_md", color="muted",
             size_px=14)
        y = 210
    hf, hc = ("primary", "on_primary") if header_color == "coral" else \
             ("surface_dark", "on_dark")
    n = len(data[0])
    add_table(s, LX, y, CW, min(360, 64 * len(data)), data,
              col_widths=[CW / n] * n, header_fill=hf, header_color=hc,
              highlight_col=highlight_col, label_col=label_col)
    deck.footer(s)
    return s


def process(deck, title, steps, eyebrow="Process"):
    """가로 프로세스. steps: [(title, desc[, color])]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(steps)
    gap = 28
    cw = (CW - (n - 1) * gap) / n
    y = 280
    line_div(s, LX + cw / 2, y + 28, CW - cw, color="hairline", weight=2)
    for i, st in enumerate(steps):
        t, desc = st[0], st[1]
        color = _accent(i, st[2] if len(st) > 2 else None)
        x = LX + i * (cw + gap)
        oval(s, x + cw / 2 - 28, y, 56, 56, fill=color)
        text(s, x + cw / 2 - 28, y, 56, 56, str(i + 1), type_token="title_lg",
             color="on_dark" if color == "surface_dark" else "on_primary",
             align="center", valign="middle")
        text(s, x, y + 76, cw, 48, t, type_token="title_sm", color="ink",
             align="center", line_spacing=1.2)
        if desc:
            text(s, x, y + 128, cw, 80, desc, type_token="body_sm",
                 color="muted", align="center", line_spacing=1.4)
    deck.footer(s)
    return s


def vsteps(deck, title, steps, eyebrow="Process"):
    """세로 단계. steps: [(title, desc[, color])]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    y = 200
    vline(s, LX + 22, y + 22, (len(steps) - 1) * 100, color="hairline",
          weight=2)
    for i, st in enumerate(steps):
        t, desc = st[0], st[1]
        color = _accent(i, st[2] if len(st) > 2 else None)
        oval(s, LX, y, 44, 44, fill=color)
        text(s, LX, y, 44, 44, str(i + 1), type_token="title_md",
             color="on_dark" if color == "surface_dark" else "on_primary",
             align="center", valign="middle")
        text(s, LX + 76, y - 2, 400, 32, t, type_token="title_lg", color="ink")
        text(s, LX + 76, y + 30, 900, 30, desc, type_token="body_md",
             color="muted")
        y += 100
    deck.footer(s)
    return s


def timeline(deck, title, nodes, eyebrow="Timeline"):
    """가로 타임라인. nodes: [(label, title, sub)].
    기간 칩 + 헤일로 컬러 노드 + 하단 정렬(라벨/설명)."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(nodes)
    line_y = 358
    x0 = LX + 70
    span = CW - 140
    step = span / (n - 1) if n > 1 else span
    line_div(s, x0, line_y, span, color="hairline", weight=3)
    colw = (span / max(1, n - 1)) * 0.92 if n > 1 else span
    for i, (lab, t, sub) in enumerate(nodes):
        x = x0 + i * step
        color = _accent(i)
        # 기간 칩 (위)
        cw_ = max(60, 28 + len(lab) * 16)
        rect(s, x - cw_ / 2, line_y - 80, cw_, 34, fill="surface_card",
             radius=999)
        text(s, x - cw_ / 2, line_y - 80, cw_, 34, lab, type_token="caption",
             color=color, align="center", valign="middle", weight="SemiBold",
             size_px=14)
        # 노드: 캔버스 링 + 컬러 원 (선 위로 도드라지게)
        oval(s, x - 16, line_y - 16, 32, 32, fill="canvas", line="hairline",
             line_w=1.5)
        oval(s, x - 9, line_y - 9, 18, 18, fill=color)
        # 타이틀 / 설명 (아래)
        text(s, x - colw / 2, line_y + 28, colw, 30, t, type_token="title_md",
             color="ink", align="center")
        if sub:
            text(s, x - colw / 2, line_y + 62, colw, 50, sub,
                 type_token="body_sm", color="muted", align="center",
                 line_spacing=1.4)
    deck.footer(s)
    return s


# ===========================================================================
# 프레임워크 / 가이드
# ===========================================================================
def matrix_2x2(deck, title, quads, note="", eyebrow="Framework",
               axis_label=""):
    """2x2 매트릭스. quads: 4개 [(label, sub[, color])] (좌상,우상,좌하,우하)."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    cx, cy, sz, g = LX, 196, 218, 12
    pos = [(0, 0), (1, 0), (0, 1), (1, 1)]
    for i, q in enumerate(quads[:4]):
        label, sub = q[0], q[1]
        color = _accent(i, q[2] if len(q) > 2 else None)
        cxi, cyi = pos[i]
        x = cx + cxi * (sz + g)
        y = cy + cyi * (sz + g)
        rect(s, x, y, sz, sz, fill="surface_card", radius=12)
        oval(s, x + 22, y + 22, 14, 14, fill=color)
        text(s, x + 22, y + 44, sz - 44, 36, label, type_token="title_lg",
             color="ink")
        text(s, x + 22, y + 86, sz - 44, 60, sub, type_token="body_sm",
             color="muted", size_px=12)
    if axis_label:
        text(s, cx, cy + 2 * sz + g + 6, 2 * sz + g, 22, axis_label,
             type_token="caption_upper", color="muted_soft", align="center")
    if note:
        rx = cx + 2 * sz + g + 48
        rect(s, rx, 196, SW - LX - rx, 448, fill="surface_dark", radius=12)
        text(s, rx + 32, 232, 300, 30, "어떻게 읽나요?", type_token="title_md",
             color="on_dark")
        text(s, rx + 32, 282, SW - LX - rx - 64, 300, note,
             type_token="body_md", color="on_dark_soft", line_spacing=1.6)
    deck.footer(s)
    return s


def do_dont(deck, title, dos, donts, eyebrow="Guide"):
    """Do / Don't. dos·donts: [str]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    cw = (CW - 40) / 2
    for col, (head, items, color, mark) in enumerate([
            ("DO", dos, "success", "✓"),
            ("DON'T", donts, "error", "✕")]):
        x = LX + col * (cw + 40)
        rect(s, x, 200, cw, 340, fill="surface_card", radius=12)
        oval(s, x + 28, 228, 40, 40, fill=color)
        text(s, x + 28, 228, 40, 40, mark, type_token="title_md",
             color="on_primary", align="center", valign="middle")
        text(s, x + 84, 234, 300, 32, head, type_token="title_lg", color="ink",
             valign="middle", weight="ExtraBold")
        bullets(s, x + 28, 300, cw - 56, 220, items, type_token="body_md",
                color="body", marker="—", marker_color=color, gap_px=12)
    deck.footer(s)
    return s


def keyword_grid(deck, title, words, cols=3, eyebrow=""):
    """키워드 그리드. words: [str] 또는 [(label, color)]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(words)
    cw = (CW - (cols - 1) * 20) / cols
    ch = 110
    for i, w in enumerate(words):
        label, color = (w, _accent(i)) if isinstance(w, str) else \
                       (w[0], _accent(i, w[1]))
        c = i % cols
        r = i // cols
        x = LX + c * (cw + 20)
        y = 200 + r * (ch + 16)
        rect(s, x, y, cw, ch, fill="surface_card", radius=12)
        oval(s, x + 24, y + ch / 2 - 18, 36, 36, fill=color)
        text(s, x + 24, y + ch / 2 - 18, 36, 36, str(i + 1),
             type_token="title_sm", color="on_primary", align="center",
             valign="middle")
        text(s, x + 76, y, cw - 90, ch, label, type_token="title_lg",
             color="ink", valign="middle")
    deck.footer(s)
    return s


def checklist(deck, title, items, dark=True, eyebrow="Summary"):
    """요약 체크리스트."""
    s = deck.add("surface_dark" if dark else "canvas")
    tc = "on_dark" if dark else "ink"
    text(s, LX, 60, CW, 22, eyebrow.upper(), type_token="caption_upper",
         color="accent_amber" if dark else "primary")
    text(s, LX, 86, CW, 60, title, type_token="display_lg", color=tc,
         weight="Light")
    y = 200
    for it in items:
        rect(s, LX, y, 32, 32,
             fill="surface_dark_elevated" if dark else "surface_card",
             line="muted" if dark else None, radius=8)
        text(s, LX, y, 32, 32, "✓", type_token="body_md",
             color="accent_teal", align="center", valign="middle")
        text(s, LX + 52, y, 980, 32, it, type_token="title_md", color=tc,
             valign="middle", weight="Regular")
        y += 64
    deck.footer(s, dark=dark)
    return s


def worksheet(deck, title, prompts, eyebrow="Activity", action=""):
    """활동지(빈칸). prompts: [(label, hint)]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    y = 200
    for label, hint in prompts:
        text(s, LX, y, 200, 48, label, type_token="title_lg", color="ink",
             valign="middle")
        rect(s, LX + 200, y, 760, 48, fill="surface_soft", line="hairline",
             radius=8)
        text(s, LX + 216, y, 740, 48, hint, type_token="body_md",
             color="muted_soft", valign="middle")
        y += 84
    if action:
        rect(s, LX, y + 6, CW, 80, fill="primary", radius=12)
        text(s, LX + 24, y + 6, CW - 48, 80, action, type_token="title_md",
             color="on_primary", valign="middle")
    deck.footer(s)
    return s


# ===========================================================================
# 타이포 무드 (포스터 / 에디토리얼)
# ===========================================================================
def poster(deck, headline, eyebrow="", sub="", dark=True):
    """Black weight 임팩트 포스터."""
    s = deck.add("surface_dark" if dark else "primary")
    tc = "on_dark" if dark else "on_primary"
    if eyebrow:
        text(s, LX, 120, CW, 30, eyebrow.upper(), type_token="caption_upper",
             color="accent_amber" if dark else "on_primary")
    text(s, LX, 170, CW, 360, headline, type_token="display_xl", color=tc,
         weight="Black", size_px=140, line_spacing=1.0)
    rect(s, LX, 560, 64, 8, fill="primary" if dark else "canvas", radius=4)
    if sub:
        text(s, LX, 590, 900, 30, sub, type_token="title_md",
             color="on_dark_soft" if dark else "on_primary", weight="Light")
    deck.footer(s, dark=True)
    return s


def editorial(deck, headline, columns, eyebrow="Essay"):
    """Thin weight 매거진/에세이. columns: [str] (1~2개)."""
    s = deck.add("canvas")
    text(s, LX, 90, 200, 24, eyebrow.upper(), type_token="caption_upper",
         color="primary")
    text(s, LX, 130, 1040, 170, headline, type_token="display_lg", color="ink",
         weight="Thin", line_spacing=1.3, size_px=44)
    line_div(s, LX, 320, CW, color="hairline")
    cols = columns[:2]
    cw = (CW - 48) / len(cols) if cols else CW
    for i, body in enumerate(cols):
        text(s, LX + i * (cw + 48), 350, cw, 220, body, type_token="body_md",
             color="body", line_spacing=1.7)
    deck.footer(s)
    return s


# ===========================================================================
# 콘텐츠 일반
# ===========================================================================
def bullets_slide(deck, title, items, eyebrow="", marker="—"):
    """제목 + 불릿 리스트. items: [str] 또는 [(marker, str)]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    bullets(s, LX, 200, CW, 380, items, type_token="title_md", color="body",
            marker=marker, marker_color="primary", gap_px=16)
    deck.footer(s)
    return s


def two_col(deck, title, left, right, eyebrow=""):
    """2단 대비. left·right: (heading, body, fill, accent)."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    gap = 40
    cw = (CW - gap) / 2
    for i, blk in enumerate([left, right]):
        heading, body = blk[0], blk[1]
        fill = blk[2] if len(blk) > 2 else ("surface_card" if i == 0
                                            else "surface_dark")
        accent = blk[3] if len(blk) > 3 else _accent(i)
        dark = fill == "surface_dark"
        x = LX + i * (cw + gap)
        rect(s, x, 200, cw, 340, fill=fill, radius=12)
        text(s, x + 32, 232, cw - 64, 36, heading, type_token="display_sm",
             color="on_dark" if dark else "ink")
        line_div(s, x + 32, 286, cw - 64,
                 color="surface_dark_elevated" if dark else "hairline")
        if isinstance(body, (list, tuple)):
            bullets(s, x + 32, 308, cw - 64, 220, body, type_token="body_md",
                    color="on_dark_soft" if dark else "body",
                    marker_color=accent, gap_px=10)
        else:
            text(s, x + 32, 308, cw - 64, 220, body, type_token="body_md",
                 color="on_dark_soft" if dark else "body", line_spacing=1.6)
    deck.footer(s)
    return s


# ===========================================================================
# 차트 (네이티브 · 편집 가능)
# ===========================================================================
def chart(deck, title, categories, series, kind="column", eyebrow="",
          note=""):
    """네이티브 차트. series: [(name, [values])]. kind: column|line|donut."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    rect(s, LX, 196, CW, 392, fill="canvas", line="hairline", radius=12)
    data = CategoryChartData()
    data.categories = categories
    for name, vals in series:
        data.add_series(name, vals)
    ctype = {"column": XL_CHART_TYPE.COLUMN_CLUSTERED,
             "line": XL_CHART_TYPE.LINE_MARKERS,
             "donut": XL_CHART_TYPE.DOUGHNUT}[kind]
    gf = s.shapes.add_chart(ctype, px(LX + 24), px(218), px(CW - 48), px(348),
                            data)
    ch = gf.chart
    ch.font.name = T.FONT_BODY
    ch.font.size = pt(20)
    ch.font.color.rgb = rgb("muted")
    palette = ["primary", "accent_teal", "accent_amber", "surface_dark"]
    if kind == "donut":
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.RIGHT
        ch.legend.include_in_layout = False
        ch.has_title = False
        for pt_, col in zip(ch.plots[0].series[0].points, palette):
            pt_.format.fill.solid()
            pt_.format.fill.fore_color.rgb = rgb(col)
    else:
        ch.has_legend = len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.TOP
            ch.legend.include_in_layout = False
        for sr, col in zip(ch.series, palette):
            if kind == "line":
                sr.format.line.color.rgb = rgb(col)
                sr.format.line.width = pt(2.5)
                sr.marker.format.fill.solid()
                sr.marker.format.fill.fore_color.rgb = rgb(col)
                sr.marker.format.line.color.rgb = rgb(col)
            else:
                sr.format.fill.solid()
                sr.format.fill.fore_color.rgb = rgb(col)
    if note:
        text(s, LX, 600, CW, 20, note, type_token="body_sm", color="muted_soft")
    deck.footer(s)
    return s


# ===========================================================================
# 가이드북 / 이미지 (스크린샷 · 단계 안내)
# ===========================================================================
def _tip_band(slide, x, y, w, body, label="TIP", h=92):
    """좌측 강조바 + 라벨 + 본문의 콜아웃 밴드 (TIP/주의/안내)."""
    rect(slide, x, y, w, h, fill="surface_cream_strong", radius=10)
    rect(slide, x, y, 6, h, fill="primary")
    text(slide, x + 28, y + 16, w - 56, 20, "▶  " + label.upper(),
         type_token="caption_upper", color="primary")
    text(slide, x + 28, y + 40, w - 56, h - 48, body, type_token="body_md",
         color="body", line_spacing=1.4)


def split_image(deck, title, url="", body="", chips=None, image=None,
                number=None, eyebrow="", title2="", title2_color="primary",
                caption="스크린샷"):
    """좌측 텍스트(번호·제목·URL·본문·칩) + 우측 스크린샷. 가이드북 핵심 레이아웃.
    image=파일경로(없으면 플레이스홀더). number를 주면 좌측 큰 번호 배지."""
    s = deck.add("canvas")
    if eyebrow:
        text(s, LX, 70, 600, 22, eyebrow.upper(), type_token="caption_upper",
             color="muted")
    tx, ty = LX, 244
    if number is not None:
        oval(s, LX, ty + 4, 104, 104, fill="primary")
        text(s, LX, ty + 4, 104, 104, str(number), type_token="display_lg",
             color="on_primary", align="center", valign="middle",
             weight="ExtraBold", size_px=46)
        tx = LX + 132
    if title2:
        text(s, tx, ty, 520, 70, title, type_token="display_lg", color="ink",
             weight="ExtraBold", size_px=50)
        text(s, tx, ty + 66, 520, 70, title2, type_token="display_lg",
             color=title2_color, weight="ExtraBold", size_px=50)
        uy = ty + 150
    else:
        text(s, tx, ty + 16, 520, 80, title, type_token="display_lg",
             color="ink", weight="ExtraBold", size_px=50)
        uy = ty + 104
    if url:
        text(s, tx, uy, 560, 30, url, type_token="title_md", color="primary",
             weight="Regular")
        uy += 50
    if body:
        text(s, LX, max(uy, 466), 560, 110, body, type_token="body_md",
             color="body", line_spacing=1.6)
    if chips:
        cx = LX
        for c in chips:
            cx += _chip(s, cx, 600, c) + 10
    picture(s, 672, 150, 528, 432, path=image, caption=caption)
    deck.footer(s)
    return s


def steps_image(deck, title, steps, image=None, eyebrow="", caption="스크린샷"):
    """좌측 번호 단계(제목+설명) + 우측 스크린샷. 가이드북 'How to' 레이아웃.
    steps: [(title, desc[, color])] (권장 3~5)."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    y = 214
    gap = 92 if len(steps) <= 5 else 76
    lw = 540
    for i, st in enumerate(steps):
        t, desc = st[0], st[1]
        color = _accent(i, st[2] if len(st) > 2 else None)
        oval(s, LX, y, 38, 38, fill=color)
        text(s, LX, y, 38, 38, str(i + 1), type_token="title_sm",
             color="on_primary", align="center", valign="middle")
        text(s, LX + 56, y - 4, lw, 32, t, type_token="title_md", color="ink",
             weight="SemiBold")
        if desc:
            text(s, LX + 56, y + 30, lw, 50, desc, type_token="body_sm",
                 color="muted", line_spacing=1.4)
        y += gap
    picture(s, 700, 170, 500, 430, path=image, caption=caption)
    deck.footer(s)
    return s


def index_list(deck, title, items, eyebrow="", sub=""):
    """좌측 타이틀 + 우측 번호 디렉터리(제목 + URL/보조). 목차·도구 목록용.
    items: [(title, url/sub[, color])] (권장 3~6)."""
    s = deck.add("canvas")
    if eyebrow:
        text(s, LX, 150, 460, 22, eyebrow.upper(), type_token="caption_upper",
             color="muted")
    text(s, LX, 186, 470, 160, title, type_token="display_xl", color="ink",
         weight="ExtraBold", size_px=72, line_spacing=1.05)
    if sub:
        sub_y = 372 if "\n" in title else 320
        text(s, LX, sub_y, 460, 40, sub, type_token="title_md", color="muted")
    rx = 600
    n = len(items)
    avail = 540
    gap = min(96, avail / n)
    y = 130 + (avail - gap * n) / 2 if n < 6 else 120
    for i, it in enumerate(items):
        t, su = it[0], it[1]
        color = _accent(i, it[2] if len(it) > 2 else None)
        oval(s, rx, y, 40, 40, fill=color)
        text(s, rx, y, 40, 40, str(i + 1), type_token="title_sm",
             color="on_primary", align="center", valign="middle")
        text(s, rx + 60, y - 6, 600, 32, t, type_token="title_lg", color="ink",
             weight="SemiBold")
        if su:
            text(s, rx + 60, y + 28, 600, 26, su, type_token="body_sm",
                 color="primary")
        y += gap
    deck.footer(s)
    return s


def feature_cards(deck, title, cards, cols=None, tip="", tip_label="TIP",
                  eyebrow=""):
    """N-up 카드(3~6, three_up 일반화). 하단 TIP 밴드 선택.
    cards: [(title, body)] 또는 [(badge, title, body[, color])].
    badge(글자/숫자)가 있으면 카드 상단에 정사각 배지 + 제목 한 줄."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(cards)
    if cols is None:
        cols = n if n <= 5 else 3
    rows = (n + cols - 1) // cols
    gap = 20
    cw = (CW - (cols - 1) * gap) / cols
    if rows > 1:                      # 2행 — 슬라이드/푸터 안에 맞춤
        top, ch, rh = 178, 232, 248
    elif tip:                        # 1행 + 하단 TIP 밴드
        top, ch, rh = 196, 248, 268
    else:                            # 1행
        top, ch, rh = 196, 300, 320
    for i, c in enumerate(cards):
        if len(c) >= 3:
            badge, t, body = c[0], c[1], c[2]
            color = _accent(i, c[3] if len(c) > 3 else None)
        else:
            badge, t, body = None, c[0], c[1]
            color = _accent(i)
        r, cc = i // cols, i % cols
        x = LX + cc * (cw + gap)
        y = top + r * rh
        rect(s, x, y, cw, ch, fill="surface_card", radius=12)
        if badge:
            rect(s, x + 24, y + 24, 36, 36, fill=color, radius=8)
            text(s, x + 24, y + 24, 36, 36, str(badge), type_token="title_sm",
                 color="on_primary", align="center", valign="middle")
            text(s, x + 72, y + 24, cw - 96, 36, t, type_token="title_md",
                 color="ink", valign="middle", weight="SemiBold")
            byo = y + 80
        else:
            oval(s, x + 28, y + 24, 42, 42, fill=color)
            text(s, x + 28, y + 84, cw - 56, 38, t, type_token="title_lg",
                 color="ink")
            byo = y + 128
        text(s, x + 24, byo, cw - 48, ch - (byo - y) - 16, body,
             type_token="body_sm", color="body", line_spacing=1.45)
    if tip:
        _tip_band(s, LX, top + rows * rh + 4, CW, tip, label=tip_label)
    deck.footer(s)
    return s


def check_cards(deck, title, items, eyebrow="", intro="", tip="",
                tip_label="TIP"):
    """좌측 타이틀(+안내/TIP) + 우측 번호 체크 카드. 최종 점검·체크리스트용.
    items: [(title, desc)] (권장 4~6)."""
    s = deck.add("canvas")
    if eyebrow:
        text(s, LX, 110, 460, 22, eyebrow.upper(), type_token="caption_upper",
             color="muted")
    text(s, LX, 146, 470, 170, title, type_token="display_xl", color="ink",
         weight="ExtraBold", size_px=66, line_spacing=1.05)
    if intro:
        text(s, LX, 330, 460, 80, intro, type_token="body_md", color="muted",
             line_spacing=1.6)
    if tip:
        _tip_band(s, LX, 446, 470, tip, label=tip_label)
    rx = 600
    rw = SW - LX - rx
    n = len(items)
    avail = 530
    chh = min(86, (avail - (n - 1) * 12) / n)
    y = 110 + (avail - (chh * n + 12 * (n - 1))) / 2
    for i, it in enumerate(items):
        t, desc = it[0], it[1]
        rect(s, rx, y, rw, chh, fill="surface_card", radius=10)
        rect(s, rx + 22, y + chh / 2 - 16, 32, 32, fill="surface_soft",
             line="hairline", radius=7)
        text(s, rx + 22, y + chh / 2 - 16, 32, 32, "✓", color="accent_teal",
             align="center", valign="middle", type_token="title_sm")
        text(s, rx + 72, y + 14, rw - 160, 28, t, type_token="title_md",
             color="ink", weight="SemiBold")
        if desc:
            text(s, rx + 72, y + 46, rw - 160, 26, desc, type_token="body_sm",
                 color="muted")
        oval(s, rx + rw - 52, y + chh / 2 - 16, 32, 32, fill="surface_soft")
        text(s, rx + rw - 52, y + chh / 2 - 16, 32, 32, str(i + 1),
             type_token="title_sm", color="primary", align="center",
             valign="middle")
        y += chh + 12
    deck.footer(s)
    return s


# ===========================================================================
# 마무리
# ===========================================================================
def closing(deck, title, sub="", contact="", eyebrow="", action=""):
    """풀블리드 다크 CTA 클로징 — 중앙 정렬, 슬림 버튼. 모든 테마에서 프리미엄.
    action: 버튼 라벨(예: '문의하기'). 주면 버튼=action, 그 아래 contact를 작게 표기."""
    s = deck.add("surface_dark")
    # 상단 짧은 액센트
    rect(s, SW / 2 - 24, 196, 48, 4, fill="primary", radius=2)
    if eyebrow:
        text(s, LX, 220, CW, 22, eyebrow.upper(), type_token="caption_upper",
             color="primary", align="center")
    text(s, LX, 262, CW, 100, title, type_token="display_lg",
         color="on_dark", align="center", weight="Light")
    y = 372
    if sub:
        text(s, LX, y, CW, 40, sub, type_token="display_sm",
             color="on_dark_soft", align="center", weight="Light")
        y += 70
    else:
        y = 412
    if action or contact:
        label = action or contact
        bw = max(220, 64 + len(label) * 13)
        bx = SW / 2 - bw / 2
        rect(s, bx, y, bw, 54, fill="primary", radius=8)
        text(s, bx, y, bw, 54, label, type_token="title_sm",
             color="on_primary", align="center", valign="middle")
        if action and contact:
            text(s, LX, y + 66, CW, 24, contact, type_token="body_md",
                 color="on_dark_soft", align="center", size_px=15)
    deck.footer(s, dark=True)
    return s


# ===========================================================================
# 범용 — 제안서 · 보고서 · 강의 공통 (목차 · KPI · 가격 · 팀 · 로드맵 등)
# ===========================================================================
def agenda(deck, title, items, eyebrow="Agenda", active=None):
    """목차/아젠다. 큰 번호 + 제목 행. active 인덱스를 주면 그 항목만 강조.
    items: [str] 또는 [(제목, _)]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(items)
    y0 = 204
    rowh = min(96, 510 / n)
    for i, it in enumerate(items):
        y = y0 + i * rowh
        on = (active == i)
        ncolor = "primary" if on else "muted_soft"
        label = it if isinstance(it, str) else it[0]
        text(s, LX, y, 92, rowh, f"{i+1:02d}", type_token="display_md",
             color=ncolor, weight="ExtraBold", valign="middle", size_px=40)
        text(s, LX + 124, y, CW - 124, rowh, label, type_token="title_lg",
             color="ink" if on else "body", valign="middle",
             weight="SemiBold" if on else "Regular")
        if i < n - 1:
            line_div(s, LX, y + rowh, CW, color="hairline")
    deck.footer(s)
    return s


def kpi_row(deck, title, kpis, eyebrow=""):
    """가로 KPI 행(2~5). 카드 없이 구분선. kpis: [(값, 라벨[, 색])]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = min(len(kpis), 5)
    cw = CW / n
    y = 286
    for i in range(n):
        k = kpis[i]
        val, label = k[0], k[1]
        color = _accent(i, k[2] if len(k) > 2 else None)
        x = LX + i * cw
        if i > 0:
            vline(s, x, y, 150, color="hairline", weight=1.5)
        text(s, x + 28, y, cw - 48, 90, val, type_token="display_lg",
             color=color, weight="ExtraBold", size_px=62)
        text(s, x + 28, y + 110, cw - 48, 70, label, type_token="body_md",
             color="muted", line_spacing=1.4)
    deck.footer(s)
    return s


def pricing(deck, title, plans, eyebrow="", highlight=None):
    """가격 플랜(최대 3). plans: [(이름, 가격, 부가, [기능들])].
    highlight 인덱스는 다크 카드 + 강조 헤더."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = min(len(plans), 3)
    gap = 24
    cw = (CW - (n - 1) * gap) / n
    for i in range(n):
        name, price, sub, feats = plans[i][0], plans[i][1], plans[i][2], \
            plans[i][3]
        on = (highlight == i)
        x = LX + i * (cw + gap)
        fill = "surface_dark" if on else "surface_card"
        tc = "on_dark" if on else "ink"
        rect(s, x, 190, cw, 404, fill=fill, radius=14)
        if on:
            bw = 92
            rect(s, x + cw - bw - 24, 210, bw, 30, fill="primary", radius=999)
            text(s, x + cw - bw - 24, 210, bw, 30, "추천",
                 type_token="caption", color="on_primary", align="center",
                 valign="middle", size_px=13)
        text(s, x + 30, 224, cw - 60, 30, name, type_token="title_md",
             color="accent_amber" if on else "primary", weight="SemiBold")
        text(s, x + 30, 262, cw - 60, 64, price, type_token="display_lg",
             color=tc, weight="ExtraBold", size_px=50)
        if sub:
            text(s, x + 30, 332, cw - 60, 24, sub, type_token="body_sm",
                 color="on_dark_soft" if on else "muted")
        line_div(s, x + 30, 368, cw - 60,
                 color="surface_dark_elevated" if on else "hairline")
        bullets(s, x + 30, 388, cw - 60, 190, feats, type_token="body_sm",
                color="on_dark_soft" if on else "body", marker="✓",
                marker_color="accent_teal", gap_px=11)
    deck.footer(s)
    return s


def team(deck, title, people, eyebrow=""):
    """팀/인물 카드(최대 4). people: [(이름, 역할[, 한줄소개])]. 이니셜 아바타."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = min(len(people), 4)
    gap = 24
    cw = (CW - (n - 1) * gap) / n
    for i in range(n):
        p = people[i]
        name, role = p[0], p[1]
        desc = p[2] if len(p) > 2 else ""
        color = _accent(i)
        x = LX + i * (cw + gap)
        rect(s, x, 200, cw, 326, fill="surface_card", radius=12)
        oval(s, x + cw / 2 - 44, 232, 88, 88, fill=color)
        text(s, x + cw / 2 - 44, 232, 88, 88, name[:2], type_token="display_sm",
             color="on_primary", align="center", valign="middle",
             weight="ExtraBold")
        text(s, x + 16, 342, cw - 32, 30, name, type_token="title_lg",
             color="ink", align="center")
        text(s, x + 16, 382, cw - 32, 24, role, type_token="body_sm",
             color="primary", align="center", weight="SemiBold")
        if desc:
            text(s, x + 16, 414, cw - 32, 96, desc, type_token="body_sm",
                 color="muted", align="center", line_spacing=1.45)
    deck.footer(s)
    return s


def logos(deck, title, names, cols=4, eyebrow="", note=""):
    """로고/고객사·도입처 그리드(텍스트 셀). names: [str]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(names)
    rows = (n + cols - 1) // cols
    gap = 20
    cw = (CW - (cols - 1) * gap) / cols
    ch = 96
    top = 214
    for i, nm in enumerate(names):
        r, c = i // cols, i % cols
        x = LX + c * (cw + gap)
        y = top + r * (ch + gap)
        rect(s, x, y, cw, ch, fill="surface_card", radius=10)
        text(s, x, y, cw, ch, nm, type_token="title_md", color="muted",
             align="center", valign="middle", weight="SemiBold")
    if note:
        text(s, LX, top + rows * (ch + gap) + 8, CW, 24, note,
             type_token="body_sm", color="muted_soft", align="center")
    deck.footer(s)
    return s


def roadmap(deck, title, phases, eyebrow="Roadmap"):
    """단계/분기 로드맵 스윔레인(최대 4). phases: [(라벨, 제목, [항목들])]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = min(len(phases), 4)
    gap = 20
    cw = (CW - (n - 1) * gap) / n
    top = 200
    for i in range(n):
        lab, t, items = phases[i][0], phases[i][1], phases[i][2]
        color = _accent(i)
        x = LX + i * (cw + gap)
        rect(s, x, top, cw, 60, fill=color, radius=10)
        text(s, x + 22, top, cw - 44, 60, lab, type_token="title_md",
             color="on_dark" if color == "surface_dark" else "on_primary",
             valign="middle", weight="SemiBold")
        rect(s, x, top + 70, cw, 290, fill="surface_card", radius=10)
        text(s, x + 22, top + 90, cw - 44, 30, t, type_token="title_md",
             color="ink", weight="SemiBold")
        bullets(s, x + 22, top + 132, cw - 44, 210, items, type_token="body_sm",
                color="body", marker="—", marker_color=color, gap_px=10)
    deck.footer(s)
    return s


def callout(deck, headline, body="", eyebrow="", tone="primary", icon="!"):
    """단독 강조 슬라이드(풀블리드 컬러/다크). 전환·경고·핵심 한 마디.
    tone: 'primary' | 'dark' | 'accent_teal' 등 색 토큰."""
    dark = tone == "dark"
    s = deck.add("surface_dark" if dark else tone)
    on = "on_dark" if dark else "on_primary"
    soft = "on_dark_soft" if dark else "on_primary"
    oval(s, LX, 150, 72, 72, fill="canvas")
    text(s, LX, 150, 72, 72, icon, type_token="display_sm",
         color="primary" if dark else tone, align="center", valign="middle",
         weight="Black")
    if eyebrow:
        text(s, LX, 258, CW, 24, eyebrow.upper(), type_token="caption_upper",
             color=soft)
    text(s, LX, 294, CW, 200, headline, type_token="display_lg", color=on,
         weight="Light", line_spacing=1.2)
    if body:
        text(s, LX, 502, 960, 90, body, type_token="title_md", color=soft,
             weight="Light", line_spacing=1.5)
    deck.footer(s, dark=True)
    return s


def metric_callout(deck, number, title, points, pre="", eyebrow=""):
    """좌측 거대한 숫자 + 우측 제목·불릿 설명 패널. 임팩트 + 근거 동시 제시.
    points: [str]."""
    s = deck.add("canvas")
    if eyebrow:
        text(s, LX, 96, 520, 24, eyebrow.upper(), type_token="caption_upper",
             color="primary")
    if pre:
        text(s, LX, 154, 520, 30, pre, type_token="title_md", color="muted")
    text(s, LX, 192, 520, 220, number, type_token="display_xl", color="primary",
         weight="Black", size_px=168)
    rx = 640
    rect(s, rx, 150, SW - LX - rx, 440, fill="surface_card", radius=14)
    text(s, rx + 36, 188, SW - LX - rx - 72, 90, title, type_token="display_sm",
         color="ink", weight="SemiBold", line_spacing=1.25)
    bullets(s, rx + 36, 300, SW - LX - rx - 72, 260, points,
            type_token="body_md", color="body", marker="—",
            marker_color="primary", gap_px=14)
    deck.footer(s)
    return s


def feature_list(deck, title, features, eyebrow=""):
    """2열 아이콘+제목+설명 리스트(최대 6). 카드보다 조밀.
    features: [(배지, 제목, 설명)] 또는 [(제목, 설명)]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = min(len(features), 6)
    gap_x = 48
    cw = (CW - gap_x) / 2
    rows = (n + 1) // 2
    top = 204
    rh = min(124, 392 / rows)
    for i in range(n):
        f = features[i]
        if len(f) >= 3:
            badge, t, desc = f[0], f[1], f[2]
        else:
            badge, t, desc = str(i + 1), f[0], f[1]
        color = _accent(i)
        r, c = i // 2, i % 2
        x = LX + c * (cw + gap_x)
        y = top + r * rh
        oval(s, x, y, 48, 48, fill=color)
        text(s, x, y, 48, 48, str(badge), type_token="title_md",
             color="on_primary", align="center", valign="middle")
        text(s, x + 68, y - 2, cw - 68, 30, t, type_token="title_md",
             color="ink", weight="SemiBold")
        text(s, x + 68, y + 32, cw - 68, 64, desc, type_token="body_sm",
             color="muted", line_spacing=1.45)
    deck.footer(s)
    return s


def progress_bars(deck, title, bars, eyebrow=""):
    """가로 진행률/역량 바 리스트. bars: [(라벨, 퍼센트[, 색])]."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(bars)
    y = 212
    rowh = min(90, 420 / n)
    for i, b in enumerate(bars):
        label, pct = b[0], b[1]
        color = _accent(i, b[2] if len(b) > 2 else None)
        text(s, LX, y, 500, 26, label, type_token="title_md", color="ink",
             weight="SemiBold")
        text(s, LX, y, CW, 26, f"{pct}%", type_token="title_md", color=color,
             align="right", weight="SemiBold")
        ty = y + 38
        rect(s, LX, ty, CW, 14, fill="surface_card", radius=7)
        rect(s, LX, ty, CW * min(100, pct) / 100, 14, fill=color, radius=7)
        y += rowh
    deck.footer(s)
    return s


def swot(deck, title, quads, eyebrow="SWOT"):
    """4분면(색 헤더 + 불릿). quads: 4개 [(헤더, [항목들])] 순서=좌상,우상,좌하,우하."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    cx, cy, gap = LX, 200, 20
    cw = (CW - gap) / 2
    ch = (420 - gap) / 2
    pos = [(0, 0), (1, 0), (0, 1), (1, 1)]
    colors = ["accent_teal", "primary", "accent_amber", "error"]
    for i, q in enumerate(quads[:4]):
        head, items = q[0], q[1]
        color = colors[i]
        cc, rr = pos[i]
        x = cx + cc * (cw + gap)
        y = cy + rr * (ch + gap)
        rect(s, x, y, cw, ch, fill="surface_card", radius=12)
        rect(s, x, y, 6, ch, fill=color)
        text(s, x + 28, y + 18, cw - 56, 28, head, type_token="title_md",
             color=color, weight="ExtraBold")
        bullets(s, x + 28, y + 58, cw - 56, ch - 70, items,
                type_token="body_sm", color="body", marker="—",
                marker_color=color, gap_px=8)
    deck.footer(s)
    return s


def compare_cards(deck, title, left, right, eyebrow="", verdict=""):
    """좌우 큰 대비 카드 + 중앙 VS 배지. before/after, 경쟁 비교.
    left·right: (헤딩, [항목들][, 강조색])."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    gap = 80
    cw = (CW - gap) / 2
    for i, blk in enumerate([left, right]):
        head, items = blk[0], blk[1]
        dark = (i == 1)
        fill = "surface_dark" if dark else "surface_card"
        accent = blk[2] if len(blk) > 2 else ("muted" if i == 0 else "primary")
        x = LX + i * (cw + gap)
        rect(s, x, 200, cw, 358, fill=fill, radius=14)
        text(s, x + 32, 232, cw - 64, 34, head, type_token="display_sm",
             color="on_dark" if dark else "ink", weight="SemiBold")
        line_div(s, x + 32, 288, cw - 64,
                 color="surface_dark_elevated" if dark else "hairline")
        bullets(s, x + 32, 308, cw - 64, 230, items, type_token="body_md",
                color="on_dark_soft" if dark else "body", marker="—",
                marker_color=accent, gap_px=11)
    oval(s, SW / 2 - 34, 344, 68, 68, fill="primary")
    text(s, SW / 2 - 34, 344, 68, 68, "VS", type_token="title_md",
         color="on_primary", align="center", valign="middle", weight="ExtraBold")
    if verdict:
        text(s, LX, 576, CW, 28, verdict, type_token="body_md", color="muted",
             align="center")
    deck.footer(s)
    return s


def image_full(deck, title="", caption="", image=None, eyebrow=""):
    """풀블리드 이미지 + 하단 다크 캡션 밴드. image 없으면 다크 플레이스홀더."""
    s = deck.add("surface_dark")
    picture(s, 0, 0, SW, T.SLIDE_H_PX, path=image, frame=False, radius=0,
            caption=caption or "이미지")
    if title or caption or eyebrow:
        rect(s, 0, T.SLIDE_H_PX - 156, SW, 156, fill="surface_dark")
        if eyebrow:
            text(s, LX, T.SLIDE_H_PX - 142, CW, 22, eyebrow.upper(),
                 type_token="caption_upper", color="accent_amber")
        if title:
            text(s, LX, T.SLIDE_H_PX - 112, CW, 50, title,
                 type_token="display_sm", color="on_dark", weight="SemiBold")
        if caption:
            text(s, LX, T.SLIDE_H_PX - 54, CW, 28, caption, type_token="body_md",
                 color="on_dark_soft")
    return s


def image_grid(deck, title, items, eyebrow="", cols=None):
    """이미지 갤러리 그리드 + 캡션. items: [(이미지경로, 캡션)] (없으면 플레이스홀더)."""
    s = deck.add("canvas")
    _header(s, eyebrow, title)
    n = len(items)
    cols = cols or (n if n <= 3 else (2 if n == 4 else 3))
    rows = (n + cols - 1) // cols
    gap = 20
    cw = (CW - (cols - 1) * gap) / cols
    top = 200
    ch = (392 - (rows - 1) * (gap + 30)) / rows
    for i, it in enumerate(items):
        img, cap = it[0], it[1]
        r, c = i // cols, i % cols
        x = LX + c * (cw + gap)
        y = top + r * (ch + gap + 30)
        picture(s, x, y, cw, ch, path=img, caption=cap or "이미지")
        if cap:
            text(s, x, y + ch + 6, cw, 24, cap, type_token="body_sm",
                 color="muted", align="center")
    deck.footer(s)
    return s
